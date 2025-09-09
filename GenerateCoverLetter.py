from docx import Document
import os

def generate_cover_letter(template_path, company, position, city, skill, output_root="C:\\Resume\\CoverLetter"):
    # Load template
    doc = Document(template_path)

    # Define replacements
    replacements = {
        "[company name]": company,
        "[position name]": position,
        "[City]": city,
        "[skills]":skill
    }

    # Replace placeholders
    for para in doc.paragraphs:
        for key, value in replacements.items():
            if key in para.text:
                for run in para.runs:
                    run.text = run.text.replace(key, value)

    # Build output path
    output_dir = os.path.join(output_root, company, position)
    os.makedirs(output_dir, exist_ok=True)  # creates folders if they don’t exist
    output_path = os.path.join(output_dir, "CoverLetter.docx")

    # Save
    doc.save(output_path)
    print(f"Cover letter saved to: {output_path}")

# Example usage
generate_cover_letter(
    template_path="C:\\Resume\\CoverLetter\\CoverLetterTEMPLATE3.docx",
    company="Texas Instruments",
    position="Hardware Test Engineer",
    city="Dallas",
    skill="",
)
